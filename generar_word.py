from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


# ============================================================
# GENERADOR DE INFORME WORD
# ============================================================
# Este archivo toma las tablas CSV y los gráficos PNG generados
# por main.py y arma un documento Word final.
#
# No recalcula estadísticas. Solo presenta, explica e interpreta.
# ============================================================


BASE_DIR = Path(__file__).resolve().parent
RESULTADOS_DIR = BASE_DIR / "resultados"
TABLAS_DIR = RESULTADOS_DIR / "tablas"
GRAFICOS_DIR = RESULTADOS_DIR / "graficos"
SALIDA_DIR = RESULTADOS_DIR / "informe_final"

ARCHIVO_SALIDA = SALIDA_DIR / "Informe_Final_Factor_Localia_LPF_2023.docx"


# ============================================================
# FUNCIONES AUXILIARES DE FORMATO
# ============================================================


def configurar_documento(documento):
    """
    Configura márgenes y fuente general.
    """

    section = documento.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = documento.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)


def agregar_titulo(documento, texto):
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(20)


def agregar_subtitulo(documento, texto):
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(texto)
    run.font.size = Pt(13)


def agregar_parrafo(documento, texto):
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(texto)


def agregar_formula(documento, texto):
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)

    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(11)


def agregar_bullet(documento, texto):
    p = documento.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(texto)


def agregar_imagen(documento, archivo, titulo, explicacion=None, formula=None):
    """
    Inserta un gráfico y agrega explicación debajo.
    """

    ruta = GRAFICOS_DIR / archivo

    documento.add_heading(titulo, level=3)

    if ruta.exists():
        documento.add_picture(str(ruta), width=Inches(6.3))
        ultimo = documento.paragraphs[-1]
        ultimo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        agregar_parrafo(documento, f"[No se encontró el gráfico: {archivo}]")

    if formula:
        agregar_formula(documento, formula)

    if explicacion:
        agregar_parrafo(documento, explicacion)


def formatear_valor(valor):
    """
    Formatea valores para tablas Word.
    """

    if pd.isna(valor):
        return ""

    if isinstance(valor, float):
        return f"{valor:.4f}"

    texto = str(valor)

    try:
        numero = float(texto)
        return f"{numero:.4f}"
    except ValueError:
        return texto


def agregar_tabla_dataframe(documento, df, titulo, font_size=8):
    """
    Inserta un DataFrame como tabla de Word.
    """

    documento.add_heading(titulo, level=3)

    df = df.copy()

    tabla = documento.add_table(rows=1, cols=len(df.columns))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = tabla.rows[0].cells

    for i, col in enumerate(df.columns):
        hdr_cells[i].text = str(col)
        hdr_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for parrafo in hdr_cells[i].paragraphs:
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in parrafo.runs:
                run.bold = True
                run.font.size = Pt(font_size)

    for _, fila in df.iterrows():
        row_cells = tabla.add_row().cells

        for i, valor in enumerate(fila):
            row_cells[i].text = formatear_valor(valor)
            row_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for parrafo in row_cells[i].paragraphs:
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in parrafo.runs:
                    run.font.size = Pt(font_size)

    documento.add_paragraph()


def leer_csv(nombre_archivo, **kwargs):
    """
    Lee un CSV desde resultados/tablas.
    """

    ruta = TABLAS_DIR / nombre_archivo

    if not ruta.exists():
        print(f"Advertencia: no se encontró {ruta}")
        return pd.DataFrame()

    return pd.read_csv(ruta, **kwargs)


def obtener_resultado(inferencia, bloque, medida, defecto=""):
    """
    Busca un valor dentro del CSV resultados_inferencia.csv.
    """

    if inferencia.empty:
        return defecto

    filtro = (inferencia["bloque"] == bloque) & (inferencia["medida"] == medida)

    if filtro.any():
        return inferencia.loc[filtro, "valor"].iloc[0]

    return defecto


def pvalor_legible(valor):
    """
    Convierte p-valores muy chicos a formato legible.
    """

    try:
        v = float(valor)
    except (ValueError, TypeError):
        return str(valor)

    if v < 0.0001:
        return "< 0.0001"

    return f"{v:.6f}"


def numero_legible(valor, decimales=4):
    """
    Convierte números a texto redondeado.
    """

    try:
        return f"{float(valor):.{decimales}f}"
    except (ValueError, TypeError):
        return str(valor)


# ============================================================
# TABLAS INTERPRETADAS
# ============================================================


def crear_tabla_inferencia_interpretada(inferencia):
    """
    Crea una tabla clara para el Word.
    Evita mostrar el CSV crudo bloque-medida-valor.
    """

    ic_inf = obtener_resultado(
        inferencia,
        "intervalo_confianza_delta_puntos",
        "ic_inferior",
    )

    ic_sup = obtener_resultado(
        inferencia,
        "intervalo_confianza_delta_puntos",
        "ic_superior",
    )

    t_puntos = obtener_resultado(inferencia, "prueba_t_puntos", "estadistico")
    p_t_puntos = obtener_resultado(inferencia, "prueba_t_puntos", "p_valor")
    dec_t_puntos = obtener_resultado(inferencia, "prueba_t_puntos", "decision")

    t_goles = obtener_resultado(inferencia, "prueba_t_goles", "estadistico")
    p_t_goles = obtener_resultado(inferencia, "prueba_t_goles", "p_valor")
    dec_t_goles = obtener_resultado(inferencia, "prueba_t_goles", "decision")

    w = obtener_resultado(inferencia, "wilcoxon_puntos", "estadistico")
    p_w = obtener_resultado(inferencia, "wilcoxon_puntos", "p_valor")
    dec_w = obtener_resultado(inferencia, "wilcoxon_puntos", "decision")

    positivos = obtener_resultado(inferencia, "prueba_signo_puntos", "positivos")
    negativos = obtener_resultado(inferencia, "prueba_signo_puntos", "negativos")
    p_signo = obtener_resultado(inferencia, "prueba_signo_puntos", "p_valor")
    dec_signo = obtener_resultado(inferencia, "prueba_signo_puntos", "decision")

    vict_local = obtener_resultado(
        inferencia,
        "prueba_binomial_victorias",
        "victorias_local",
    )

    vict_visit = obtener_resultado(
        inferencia,
        "prueba_binomial_victorias",
        "victorias_visitante",
    )

    empates_excluidos = obtener_resultado(
        inferencia,
        "prueba_binomial_victorias",
        "empates_excluidos",
    )

    p_binomial = obtener_resultado(
        inferencia,
        "prueba_binomial_victorias",
        "p_valor",
    )

    dec_binomial = obtener_resultado(
        inferencia,
        "prueba_binomial_victorias",
        "decision",
    )

    chi2 = obtener_resultado(inferencia, "chi_cuadrado_resultados", "chi2")
    gl_chi = obtener_resultado(inferencia, "chi_cuadrado_resultados", "gl")
    p_chi = obtener_resultado(inferencia, "chi_cuadrado_resultados", "p_valor")
    dec_chi = obtener_resultado(inferencia, "chi_cuadrado_resultados", "decision")

    chi_bondad = obtener_resultado(inferencia, "bondad_ajuste_resultados", "chi2")
    gl_bondad = obtener_resultado(inferencia, "bondad_ajuste_resultados", "gl")
    p_bondad = obtener_resultado(inferencia, "bondad_ajuste_resultados", "p_valor")
    dec_bondad = obtener_resultado(inferencia, "bondad_ajuste_resultados", "decision")

    var_muestral = obtener_resultado(
        inferencia,
        "prueba_varianza_delta",
        "varianza_muestral",
    )

    var_hip = obtener_resultado(
        inferencia,
        "prueba_varianza_delta",
        "varianza_hipotetica",
    )

    p_var = obtener_resultado(inferencia, "prueba_varianza_delta", "p_valor")
    dec_var = obtener_resultado(inferencia, "prueba_varianza_delta", "decision")

    ks = obtener_resultado(
        inferencia,
        "kolmogorov_smirnov_delta",
        "estadistico",
    )

    p_ks = obtener_resultado(
        inferencia,
        "kolmogorov_smirnov_delta",
        "p_valor",
    )

    dec_ks = obtener_resultado(
        inferencia,
        "kolmogorov_smirnov_delta",
        "decision",
    )

    tabla = pd.DataFrame(
        [
            {
                "Análisis": "IC 95% para Delta Pts/PJ",
                "Hipótesis nula": "La diferencia media podría ser 0",
                "Resultado": f"[{numero_legible(ic_inf)} ; {numero_legible(ic_sup)}]",
                "p-valor": "-",
                "Decisión": "Ventaja positiva",
                "Interpretación": "El intervalo completo queda por encima de 0, por lo que la ventaja promedio de localía es positiva.",
            },
            {
                "Análisis": "Prueba t pareada - puntos",
                "Hipótesis nula": "mu_d = 0",
                "Resultado": f"t = {numero_legible(t_puntos)}",
                "p-valor": pvalor_legible(p_t_puntos),
                "Decisión": dec_t_puntos,
                "Interpretación": "Los equipos obtienen significativamente más puntos por partido como locales.",
            },
            {
                "Análisis": "Prueba t pareada - goles",
                "Hipótesis nula": "mu_d = 0",
                "Resultado": f"t = {numero_legible(t_goles)}",
                "p-valor": pvalor_legible(p_t_goles),
                "Decisión": dec_t_goles,
                "Interpretación": "Los equipos convierten significativamente más goles por partido como locales.",
            },
            {
                "Análisis": "Wilcoxon pareada",
                "Hipótesis nula": "Mediana de diferencias = 0",
                "Resultado": f"W = {numero_legible(w)}",
                "p-valor": pvalor_legible(p_w),
                "Decisión": dec_w,
                "Interpretación": "Confirma la ventaja local sin depender de normalidad estricta.",
            },
            {
                "Análisis": "Prueba del signo",
                "Hipótesis nula": "p = 0.5",
                "Resultado": f"{positivos} positivos vs {negativos} negativos",
                "p-valor": pvalor_legible(p_signo),
                "Decisión": dec_signo,
                "Interpretación": "Evalúa si la mayoría de los equipos rinde mejor como local, sin medir magnitud.",
            },
            {
                "Análisis": "Prueba binomial de victorias",
                "Hipótesis nula": "p = 0.5",
                "Resultado": f"{vict_local} victorias locales vs {vict_visit} visitantes",
                "p-valor": pvalor_legible(p_binomial),
                "Decisión": dec_binomial,
                "Interpretación": f"Los empates ({empates_excluidos}) se excluyen porque la binomial compara solo dos categorías.",
            },
            {
                "Análisis": "Chi-cuadrado r x c",
                "Hipótesis nula": "Resultado independiente de localía",
                "Resultado": f"chi2 = {numero_legible(chi2)}; gl = {gl_chi}",
                "p-valor": pvalor_legible(p_chi),
                "Decisión": dec_chi,
                "Interpretación": "El resultado del partido está asociado a jugar de local o visitante.",
            },
            {
                "Análisis": "Bondad de ajuste",
                "Hipótesis nula": "Victoria local, empate y victoria visitante son igualmente probables",
                "Resultado": f"chi2 = {numero_legible(chi_bondad)}; gl = {gl_bondad}",
                "p-valor": pvalor_legible(p_bondad),
                "Decisión": dec_bondad,
                "Interpretación": "La distribución global de resultados no se comporta como uniforme.",
            },
            {
                "Análisis": "Hipótesis sobre varianza",
                "Hipótesis nula": f"sigma2 = {numero_legible(var_hip)}",
                "Resultado": f"s2 = {numero_legible(var_muestral)}",
                "p-valor": pvalor_legible(p_var),
                "Decisión": dec_var,
                "Interpretación": "Analiza si la dispersión de la ventaja de localía difiere de una varianza de referencia.",
            },
            {
                "Análisis": "Kolmogorov-Smirnov",
                "Hipótesis nula": "Delta Pts/PJ sigue distribución normal",
                "Resultado": f"D = {numero_legible(ks)}",
                "p-valor": pvalor_legible(p_ks),
                "Decisión": dec_ks,
                "Interpretación": "Evalúa normalidad de la variable Delta Pts/PJ como complemento al Q-Q plot.",
            },
        ]
    )

    return tabla


def crear_tabla_capitulos():
    """
    Tabla que relaciona el trabajo con el programa de la materia.
    """

    return pd.DataFrame(
        [
            {
                "Capítulo": "Capítulo 1",
                "Tema": "Media, varianza, desvío, Chebyshev, percentiles",
                "Aplicación en el trabajo": "Resumen descriptivo de Delta Pts/PJ, boxplot, percentiles y Chebyshev.",
            },
            {
                "Capítulo": "Capítulo 2",
                "Tema": "Hipótesis, significancia, medias, intervalos",
                "Aplicación en el trabajo": "IC 95% y prueba t pareada para puntos y goles.",
            },
            {
                "Capítulo": "Capítulo 3",
                "Tema": "Varianza, proporciones, tablas r x c, bondad de ajuste",
                "Aplicación en el trabajo": "Prueba de varianza, binomial, chi-cuadrado y bondad de ajuste.",
            },
            {
                "Capítulo": "Capítulo 4",
                "Tema": "No paramétricas, signo, Wilcoxon, Kruskal-Wallis, KS",
                "Aplicación en el trabajo": "Prueba del signo, Wilcoxon, Kruskal-Wallis y Kolmogorov-Smirnov.",
            },
            {
                "Capítulo": "Capítulo 5",
                "Tema": "Regresión simple, mínimos cuadrados, residuos, curvilínea",
                "Aplicación en el trabajo": "Valor de plantilla vs puntos locales, residuos y ajuste polinómico.",
            },
            {
                "Capítulo": "Capítulo 6",
                "Tema": "ANOVA y comparación de grupos",
                "Aplicación en el trabajo": "Comparación de Delta Pts/PJ entre grupos económicos bajo, medio y alto.",
            },
            {
                "Capítulo": "Capítulo 7",
                "Tema": "Procesos estocásticos y estacionariedad",
                "Aplicación en el trabajo": "Se plantea como extensión futura usando varias temporadas como serie temporal.",
            },
        ]
    )


# ============================================================
# DOCUMENTO PRINCIPAL
# ============================================================


def main():
    SALIDA_DIR.mkdir(parents=True, exist_ok=True)

    documento = Document()
    configurar_documento(documento)

    # ========================================================
    # PORTADA
    # ========================================================
    agregar_titulo(documento, "Informe Final")
    agregar_subtitulo(documento, "Estadística Aplicada II")
    agregar_subtitulo(documento, "Factor localía en la Liga Profesional Argentina 2023")
    documento.add_paragraph()
    agregar_subtitulo(documento, "Análisis descriptivo, inferencial y económico")
    documento.add_paragraph()
    agregar_subtitulo(documento, "Autor: Leandro Santos")
    agregar_subtitulo(documento, "Año: 2026")

    documento.add_page_break()

    # ========================================================
    # ÍNDICE MANUAL
    # ========================================================
    documento.add_heading("Índice", level=1)

    indice = [
        "1. Introducción",
        "2. Objetivos",
        "3. Fuente de datos y metodología",
        "4. Variables, fórmulas y glosario",
        "5. Análisis descriptivo",
        "6. Percentiles, z-score y Chebyshev",
        "7. Inferencia estadística",
        "8. Empates, proporciones y tablas r x c",
        "9. Pruebas no paramétricas",
        "10. Análisis económico: valor de plantilla",
        "11. Regresión simple y ajuste de curvas",
        "12. ANOVA por grupos económicos",
        "13. Relación con los capítulos de la materia",
        "14. Conclusiones",
        "15. Limitaciones y futuras líneas de análisis",
        "16. Anexos gráficos",
    ]

    for item in indice:
        documento.add_paragraph(item)

    documento.add_page_break()

    # ========================================================
    # CARGA DE TABLAS
    # ========================================================
    tabla_global = leer_csv("resumen_global_resultados.csv")
    resumen = leer_csv("resumen_descriptivo.csv")
    ranking = leer_csv("ranking_localia.csv")
    correlaciones = leer_csv("correlaciones_plantilla.csv")
    grupos = leer_csv("resumen_grupos_economicos.csv")
    inferencia = leer_csv("resultados_inferencia.csv")
    plantilla = leer_csv("base_localia_plantilla_2023.csv")
    percentiles = leer_csv("tabla_percentiles_localia.csv")
    empates = leer_csv("resumen_empates.csv")
    formulas = leer_csv("formulas_principales.csv")
    glosario = leer_csv("glosario_estadistico.csv")
    bondad = leer_csv("bondad_ajuste_observados_esperados.csv")
    zscores = leer_csv("zscore_percentiles_localia.csv")

    # ========================================================
    # 1. INTRODUCCIÓN
    # ========================================================
    documento.add_heading("1. Introducción", level=1)

    agregar_parrafo(
        documento,
        "El presente trabajo analiza el efecto de la localía en la Liga Profesional Argentina 2023. "
        "El objetivo principal es estudiar si los equipos obtienen mejores resultados cuando juegan "
        "como locales en comparación con su rendimiento como visitantes. Además, se incorpora una "
        "dimensión económica mediante el valor de mercado de las plantillas, con el fin de evaluar "
        "si los equipos con mayor poder económico aprovechan más la condición de local.",
    )

    agregar_parrafo(
        documento,
        "El tema resulta relevante porque en el fútbol la localía suele asociarse con ventajas "
        "deportivas como el apoyo del público, el conocimiento del estadio, la reducción de viajes "
        "y factores psicológicos. Sin embargo, desde una perspectiva estadística, estas ideas deben "
        "contrastarse con datos reales, indicadores cuantitativos y pruebas formales.",
    )

    # ========================================================
    # 2. OBJETIVOS
    # ========================================================
    documento.add_heading("2. Objetivos", level=1)

    agregar_parrafo(
        documento,
        "El objetivo general es determinar si existe evidencia estadística de una ventaja de localía "
        "en la Liga Profesional Argentina 2023.",
    )

    agregar_bullet(
        documento,
        "Comparar puntos por partido, goles por partido y proporciones de resultados entre local y visitante.",
    )

    agregar_bullet(
        documento,
        "Construir e interpretar la variable Delta Pts/PJ como medida de ventaja de localía.",
    )

    agregar_bullet(
        documento,
        "Aplicar intervalos de confianza, pruebas de hipótesis, pruebas no paramétricas y análisis de proporciones.",
    )

    agregar_bullet(
        documento,
        "Incorporar el valor de plantilla para estudiar si el poder económico se relaciona con el rendimiento local.",
    )

    agregar_bullet(
        documento,
        "Vincular el análisis aplicado con los principales capítulos de Estadística Aplicada II.",
    )

    # ========================================================
    # 3. FUENTE Y METODOLOGÍA
    # ========================================================
    documento.add_heading("3. Fuente de datos y metodología", level=1)

    agregar_parrafo(
        documento,
        "La base deportiva se construyó a partir de datos de FBref correspondientes a la Liga "
        "Profesional Argentina 2023. Se utilizó la tabla Home/Away, que separa el rendimiento de "
        "cada equipo como local y como visitante. Para el análisis económico se incorporó información "
        "de Transfermarkt, utilizando el valor total de mercado de cada plantilla en millones de euros.",
    )

    agregar_parrafo(
        documento,
        "La unidad de análisis principal es cada equipo de la temporada. Para cada club se calcularon "
        "indicadores de rendimiento local y visitante. Luego se aplicaron técnicas descriptivas, "
        "inferenciales, no paramétricas, de regresión y de análisis de varianza.",
    )

    agregar_parrafo(
        documento,
        "El procesamiento fue realizado con Python. Se utilizaron pandas para carga y limpieza de datos, "
        "scipy para pruebas estadísticas y regresiones, matplotlib para gráficos y python-docx para la "
        "generación automática del informe final.",
    )

    # ========================================================
    # 4. VARIABLES, FÓRMULAS Y GLOSARIO
    # ========================================================
    documento.add_heading("4. Variables, fórmulas y glosario", level=1)

    agregar_parrafo(
        documento,
        "La variable central del trabajo es Delta Pts/PJ. Esta variable mide la diferencia entre el "
        "promedio de puntos por partido como local y el promedio de puntos por partido como visitante.",
    )

    agregar_formula(
        documento,
        "Delta Pts/PJ = Pts/PJ Local - Pts/PJ Visitante",
    )

    agregar_parrafo(
        documento,
        "Si Delta Pts/PJ es positivo, el equipo rindió mejor como local. Si es cercano a cero, el "
        "rendimiento fue similar en ambas condiciones. Si es negativo, el equipo rindió mejor como visitante.",
    )

    if not formulas.empty:
        agregar_tabla_dataframe(
            documento,
            formulas,
            "Tabla de fórmulas principales utilizadas",
            font_size=7,
        )

    if not glosario.empty:
        agregar_tabla_dataframe(
            documento,
            glosario,
            "Glosario estadístico",
            font_size=8,
        )

    # ========================================================
    # 5. ANÁLISIS DESCRIPTIVO
    # ========================================================
    documento.add_heading("5. Análisis descriptivo", level=1)

    agregar_parrafo(
        documento,
        "El análisis descriptivo muestra una diferencia clara entre el rendimiento local y visitante. "
        "Los equipos locales obtuvieron más puntos por partido, más goles por partido y una mayor "
        "proporción de victorias que los visitantes.",
    )

    if not tabla_global.empty:
        agregar_tabla_dataframe(
            documento,
            tabla_global,
            "Resumen global: local vs visitante",
            font_size=7,
        )

    agregar_imagen(
        documento,
        "01_resumen_local_visitante.png",
        "Figura 1. Resumen global local vs visitante",
        formula="Pts/PJ = Puntos totales / Partidos jugados",
        explicacion=(
            "Este gráfico compara el rendimiento global de los equipos jugando como locales y visitantes. "
            "Permite observar que los locales obtuvieron más puntos por partido, más goles por partido "
            "y mayor porcentaje de victorias. Es una primera evidencia descriptiva de la ventaja de localía."
        ),
    )

    if not ranking.empty:
        columnas_ranking = [
            "Equipo",
            "Pts_PJ_Local",
            "Pts_PJ_Visitante",
            "Delta_Pts_PJ",
            "Percentil_Localia",
            "Clasificacion_Localia",
        ]

        columnas_ranking = [c for c in columnas_ranking if c in ranking.columns]

        agregar_tabla_dataframe(
            documento,
            ranking[columnas_ranking].head(10),
            "Top 10 de equipos con mayor ventaja de localía",
            font_size=7,
        )

    agregar_parrafo(
        documento,
        "El ranking no representa necesariamente a los mejores equipos del torneo, sino a los equipos "
        "que tuvieron mayor diferencia entre su rendimiento local y visitante. Por ejemplo, un equipo "
        "puede ser muy bueno de local, pero si también fue muy bueno de visitante, su Delta Pts/PJ no será tan alto.",
    )

    agregar_imagen(
        documento,
        "02_ranking_delta_puntos.png",
        "Figura 2. Ranking de ventaja de localía",
        formula="Delta Pts/PJ = Pts/PJ Local - Pts/PJ Visitante",
        explicacion=(
            "Este gráfico ordena los equipos de mayor a menor ventaja de localía. Los primeros lugares "
            "corresponden a los equipos que más aumentaron su rendimiento jugando como locales respecto "
            "de su rendimiento visitante."
        ),
    )

    agregar_imagen(
        documento,
        "03_boxplot_delta_puntos.png",
        "Figura 3. Boxplot de puntos por partido y ventaja de localía",
        formula="RIC = Q3 - Q1",
        explicacion=(
            "El diagrama de caja y bigotes resume la distribución de Pts/PJ Local, Pts/PJ Visitante "
            "y Delta Pts/PJ. La caja muestra el rango intercuartílico, la línea central es la mediana "
            "y los bigotes muestran la dispersión principal. Se observa que los puntos por partido como "
            "local se ubican por encima de los visitantes."
        ),
    )

    # ========================================================
    # 6. PERCENTILES, Z-SCORE Y CHEBYSHEV
    # ========================================================
    documento.add_heading("6. Percentiles, z-score y Chebyshev", level=1)

    agregar_parrafo(
        documento,
        "Los percentiles permiten ubicar a cada equipo dentro de la distribución de la ventaja de localía. "
        "Por ejemplo, un equipo en percentil 90 tiene una ventaja de localía mayor o igual que aproximadamente "
        "el 90% de los equipos.",
    )

    if not percentiles.empty:
        agregar_tabla_dataframe(
            documento,
            percentiles,
            "Percentiles principales de Delta Pts/PJ",
            font_size=7,
        )

    if not zscores.empty:
        columnas_z = [
            "Equipo",
            "Delta_Pts_PJ",
            "Z_Delta_Pts_PJ",
            "Percentil_Localia",
            "Clasificacion",
        ]

        columnas_z = [c for c in columnas_z if c in zscores.columns]

        agregar_tabla_dataframe(
            documento,
            zscores[columnas_z].head(10),
            "Top 10 según z-score y percentil de localía",
            font_size=7,
        )

    agregar_imagen(
        documento,
        "06_zscore_localia.png",
        "Figura 4. Z-score de la ventaja de localía",
        formula="z = (x - media) / desvío estándar",
        explicacion=(
            "El z-score indica cuántos desvíos estándar se encuentra cada equipo respecto de la media "
            "de Delta Pts/PJ. Valores positivos indican una ventaja de localía superior al promedio. "
            "Valores negativos indican una ventaja menor al promedio."
        ),
    )

    agregar_imagen(
        documento,
        "11_chebyshev_delta_localia.png",
        "Figura 5. Teorema de Chebyshev aplicado a Delta Pts/PJ",
        formula="Al menos 1 - 1/k² de los datos cae dentro de media ± k desvíos",
        explicacion=(
            "Con k = 2, Chebyshev garantiza que al menos el 75% de los datos debe encontrarse dentro "
            "del intervalo media ± 2 desvíos estándar, sin importar la forma de la distribución. En el "
            "trabajo se compara esa cota mínima con la proporción observada."
        ),
    )

    # ========================================================
    # 7. INFERENCIA ESTADÍSTICA
    # ========================================================
    documento.add_heading("7. Inferencia estadística", level=1)

    agregar_parrafo(
        documento,
        "La inferencia estadística permite pasar de una observación descriptiva a una conclusión formal. "
        "En este caso se busca determinar si la diferencia entre rendimiento local y visitante es "
        "estadísticamente significativa.",
    )

    agregar_formula(
        documento,
        "H0: mu_d = 0     vs     H1: mu_d > 0",
    )

    agregar_parrafo(
        documento,
        "Donde d representa la diferencia Pts/PJ Local - Pts/PJ Visitante. La comparación es pareada "
        "porque cada equipo aporta dos mediciones relacionadas: una como local y otra como visitante.",
    )

    tabla_inferencia = crear_tabla_inferencia_interpretada(inferencia)

    agregar_tabla_dataframe(
        documento,
        tabla_inferencia,
        "Resumen interpretado de inferencia estadística",
        font_size=6,
    )

    agregar_imagen(
        documento,
        "10_intervalo_confianza_delta.png",
        "Figura 6. Intervalo de confianza para la ventaja media de localía",
        formula="IC = media ± t crítico * (s / raíz(n))",
        explicacion=(
            "El intervalo de confianza del 95% estima el rango plausible para la ventaja media de localía. "
            "Como el intervalo queda completamente por encima de cero, se interpreta que la ventaja promedio "
            "de localía es positiva."
        ),
    )

    agregar_imagen(
        documento,
        "12_qqplot_delta_localia.png",
        "Figura 7. Q-Q plot de Delta Pts/PJ",
        formula="Comparación de cuantiles observados vs cuantiles teóricos normales",
        explicacion=(
            "El Q-Q plot permite observar visualmente si la distribución de Delta Pts/PJ se aproxima a una "
            "normal. Se usa como complemento para justificar la aplicación de pruebas paramétricas y, a la vez, "
            "la inclusión de pruebas no paramétricas como Wilcoxon."
        ),
    )

    # ========================================================
    # 8. EMPATES, PROPORCIONES Y TABLAS R X C
    # ========================================================
    documento.add_heading("8. Empates, proporciones y tablas r x c", level=1)

    agregar_parrafo(
        documento,
        "El empate tiene un tratamiento especial. En el cálculo de puntos, el empate suma un punto para "
        "cada equipo. En la prueba binomial se excluye porque esa prueba compara solo dos categorías: "
        "victoria local contra victoria visitante. En cambio, en chi-cuadrado y bondad de ajuste el empate "
        "sí se incluye como una tercera categoría.",
    )

    if not empates.empty:
        agregar_tabla_dataframe(
            documento,
            empates,
            "Tratamiento y distribución de empates",
            font_size=7,
        )

    agregar_imagen(
        documento,
        "13_proporciones_resultados_local_visitante.png",
        "Figura 8. Proporciones de resultados: local vs visitante",
        formula="Proporción = cantidad de casos / total de partidos",
        explicacion=(
            "Este gráfico compara los porcentajes de victorias, empates y derrotas para local y visitante. "
            "Muestra que el porcentaje de victorias locales es considerablemente mayor que el porcentaje "
            "de victorias visitantes."
        ),
    )

    agregar_imagen(
        documento,
        "14_tabla_chi_cuadrado_observada.png",
        "Figura 9. Tabla observada para chi-cuadrado",
        formula="chi2 = sumatoria de (O - E)² / E",
        explicacion=(
            "La tabla r x c permite evaluar si el resultado del partido depende de la condición de local "
            "o visitante. La prueba chi-cuadrado compara frecuencias observadas contra frecuencias esperadas "
            "bajo independencia."
        ),
    )

    if not bondad.empty:
        agregar_tabla_dataframe(
            documento,
            bondad,
            "Bondad de ajuste: observados vs esperados uniformes",
            font_size=7,
        )

    # ========================================================
    # 9. PRUEBAS NO PARAMÉTRICAS
    # ========================================================
    documento.add_heading("9. Pruebas no paramétricas", level=1)

    agregar_parrafo(
        documento,
        "Las pruebas no paramétricas se utilizan como complemento cuando no se desea depender fuertemente "
        "de supuestos como la normalidad. En este trabajo se aplicaron Wilcoxon, prueba del signo, "
        "Kruskal-Wallis y Kolmogorov-Smirnov.",
    )

    agregar_bullet(
        documento,
        "Wilcoxon compara muestras relacionadas usando rangos de las diferencias.",
    )

    agregar_bullet(
        documento,
        "La prueba del signo analiza cuántos equipos tuvieron diferencia positiva y cuántos negativa.",
    )

    agregar_bullet(
        documento,
        "Kruskal-Wallis compara la distribución de la ventaja de localía entre grupos económicos.",
    )

    agregar_bullet(
        documento,
        "Kolmogorov-Smirnov evalúa si Delta Pts/PJ es compatible con una distribución normal.",
    )

    # ========================================================
    # 10. ANÁLISIS ECONÓMICO
    # ========================================================
    documento.add_heading("10. Análisis económico: valor de plantilla", level=1)

    agregar_parrafo(
        documento,
        "El análisis económico incorpora el valor total de mercado de cada plantilla. Los clubes se "
        "clasificaron en grupos económicos bajo, medio y alto mediante terciles del valor de plantilla.",
    )

    if not plantilla.empty:
        columnas_plantilla = [
            "Equipo",
            "ValorPlantillaM",
            "GrupoEconomico",
            "Pts_PJ_Local",
            "Pts_PJ_Visitante",
            "Delta_Pts_PJ",
        ]

        columnas_plantilla = [c for c in columnas_plantilla if c in plantilla.columns]

        plantilla_mostrar = plantilla[columnas_plantilla].sort_values(
            by="ValorPlantillaM",
            ascending=False,
        ).head(10)

        agregar_tabla_dataframe(
            documento,
            plantilla_mostrar,
            "Top 10 de clubes según valor de plantilla",
            font_size=7,
        )

    if not correlaciones.empty:
        agregar_tabla_dataframe(
            documento,
            correlaciones,
            "Correlaciones entre valor de plantilla y variables deportivas",
            font_size=7,
        )

    agregar_parrafo(
        documento,
        "Los resultados muestran que el valor de plantilla se asocia positivamente con los puntos por "
        "partido obtenidos como local. Sin embargo, la relación entre valor de plantilla y Delta Pts/PJ "
        "no fue significativa. Esto sugiere que los equipos con planteles más caros tienden a rendir "
        "mejor en general, pero no necesariamente tienen una ventaja de localía diferencial más alta.",
    )

    agregar_imagen(
        documento,
        "07_valor_vs_puntos_local_regresion.png",
        "Figura 10. Valor de plantilla vs puntos como local",
        formula="Pts/PJ Local = b0 + b1 * ValorPlantillaM",
        explicacion=(
            "Este gráfico representa una regresión lineal simple. La pendiente positiva indica que, "
            "en general, los equipos con mayor valor de plantilla tienden a obtener más puntos por partido "
            "como locales."
        ),
    )

    agregar_imagen(
        documento,
        "08_valor_vs_delta_localia_regresion.png",
        "Figura 11. Valor de plantilla vs ventaja de localía",
        formula="Delta Pts/PJ = b0 + b1 * ValorPlantillaM",
        explicacion=(
            "Este gráfico analiza si los equipos más caros aprovechan más la localía en términos relativos. "
            "La relación no resultó significativa, por lo que el valor de plantilla no explica claramente "
            "la diferencia entre rendimiento local y visitante."
        ),
    )

    # ========================================================
    # 11. REGRESIÓN Y AJUSTE DE CURVAS
    # ========================================================
    documento.add_heading("11. Regresión simple y ajuste de curvas", level=1)

    agregar_parrafo(
        documento,
        "La regresión simple se aplicó mediante mínimos cuadrados. Este método busca la recta que minimiza "
        "la suma de los cuadrados de los residuos, es decir, las diferencias entre los valores observados "
        "y los valores estimados por el modelo.",
    )

    agregar_formula(
        documento,
        "y = b0 + b1*x + e",
    )

    agregar_imagen(
        documento,
        "15_residuos_regresion_valor_pts_local.png",
        "Figura 12. Residuos de la regresión",
        formula="Residuo = valor observado - valor estimado",
        explicacion=(
            "El gráfico de residuos permite evaluar si el modelo lineal presenta patrones sistemáticos. "
            "Si los residuos se distribuyen alrededor de cero sin un patrón claro, el ajuste lineal es "
            "razonable como primera aproximación."
        ),
    )

    agregar_imagen(
        documento,
        "16_ajuste_polinomico_valor_pts_local.png",
        "Figura 13. Ajuste polinómico",
        formula="y = a*x² + b*x + c",
        explicacion=(
            "El ajuste polinómico se incluye como ejemplo de regresión curvilínea. Permite observar si "
            "una curva de segundo grado describe mejor la relación entre valor de plantilla y puntos como local. "
            "Su uso es complementario y no reemplaza la interpretación principal de la regresión lineal.",
        ),
    )

    # ========================================================
    # 12. ANOVA
    # ========================================================
    documento.add_heading("12. ANOVA por grupos económicos", level=1)

    agregar_parrafo(
        documento,
        "Para analizar si la ventaja de localía cambia según el grupo económico se aplicó un ANOVA de un factor. "
        "El factor es el grupo económico y la variable de respuesta es Delta Pts/PJ.",
    )

    agregar_formula(
        documento,
        "H0: media_bajo = media_medio = media_alto",
    )

    if not grupos.empty:
        agregar_tabla_dataframe(
            documento,
            grupos,
            "Resumen por grupo económico",
            font_size=7,
        )

    agregar_imagen(
        documento,
        "09_boxplot_delta_grupo_economico.png",
        "Figura 14. Boxplot de Delta Pts/PJ por grupo económico",
        formula="ANOVA compara medias entre grupos",
        explicacion=(
            "El boxplot compara la distribución de la ventaja de localía entre equipos de bajo, medio y alto "
            "valor de plantilla. Visualmente las diferencias no son marcadas, lo que coincide con el resultado "
            "no significativo del ANOVA.",
        ),
    )

    agregar_imagen(
        documento,
        "17_media_delta_por_grupo_economico.png",
        "Figura 15. Media de Delta Pts/PJ por grupo económico",
        formula="Media del grupo = suma de Delta Pts/PJ del grupo / cantidad de equipos del grupo",
        explicacion=(
            "Este gráfico muestra la media de la ventaja de localía para cada grupo económico. Aunque hay "
            "pequeñas diferencias entre grupos, las pruebas estadísticas no mostraron evidencia suficiente "
            "para afirmar que esas diferencias sean significativas.",
        ),
    )

    # ========================================================
    # 13. RELACIÓN CON CAPÍTULOS
    # ========================================================
    documento.add_heading("13. Relación con los capítulos de la materia", level=1)

    tabla_capitulos = crear_tabla_capitulos()

    agregar_tabla_dataframe(
        documento,
        tabla_capitulos,
        "Relación entre el trabajo y los capítulos del programa",
        font_size=7,
    )

    agregar_parrafo(
        documento,
        "El Capítulo 7 no se desarrolla como análisis central porque el dataset corresponde a una sola "
        "temporada. Para aplicar procesos estocásticos y estacionariedad de manera más sólida sería necesario "
        "trabajar con varias temporadas y construir una serie temporal de la ventaja de localía.",
    )

    # ========================================================
    # 14. CONCLUSIONES
    # ========================================================
    documento.add_heading("14. Conclusiones", level=1)

    agregar_parrafo(
        documento,
        "El análisis realizado permite concluir que existe una ventaja de localía significativa en la Liga "
        "Profesional Argentina 2023. Los equipos obtuvieron más puntos por partido, más goles por partido "
        "y mayor proporción de victorias cuando jugaron como locales.",
    )

    agregar_parrafo(
        documento,
        "Las pruebas estadísticas respaldan esta conclusión. El intervalo de confianza para la diferencia "
        "media de puntos por partido fue positivo; la prueba t pareada rechazó la hipótesis nula de igualdad "
        "de medias; Wilcoxon confirmó el resultado sin depender de normalidad estricta; la prueba del signo "
        "mostró mayoría de diferencias positivas; la binomial evidenció más victorias locales que visitantes; "
        "y chi-cuadrado mostró asociación entre condición de localía y resultado del partido.",
    )

    agregar_parrafo(
        documento,
        "Respecto del análisis económico, el valor de plantilla se relacionó positivamente con el rendimiento "
        "local absoluto, pero no con la ventaja de localía diferencial. Por lo tanto, los equipos más caros "
        "tienden a rendir mejor, pero no necesariamente aprovechan más la localía en términos relativos.",
    )

    # ========================================================
    # 15. LIMITACIONES
    # ========================================================
    documento.add_heading("15. Limitaciones y futuras líneas de análisis", level=1)

    agregar_parrafo(
        documento,
        "Una limitación del trabajo es que se analiza una sola temporada. Para obtener conclusiones más generales "
        "sería conveniente incorporar más años de competencia. También podrían agregarse variables como asistencia "
        "al estadio, distancia recorrida por los equipos visitantes, capacidad del estadio, altura geográfica o "
        "importancia del partido.",
    )

    agregar_parrafo(
        documento,
        "Como extensión relacionada con procesos estocásticos, podría estudiarse la evolución de la ventaja de "
        "localía a lo largo de varias temporadas y analizar si el fenómeno es estacionario o si cambia con el tiempo.",
    )

    # ========================================================
    # 16. ANEXOS
    # ========================================================
    documento.add_heading("16. Anexos gráficos", level=1)

    agregar_imagen(
        documento,
        "04_dispersion_local_visitante.png",
        "Anexo 1. Dispersión local vs visitante",
        formula="Cada punto representa un equipo",
        explicacion=(
            "La dispersión permite comparar visualmente el rendimiento local y visitante de cada equipo. "
            "Los puntos por encima de la diagonal representan equipos con mejor rendimiento local.",
        ),
    )

    agregar_imagen(
        documento,
        "05_goles_local_visitante.png",
        "Anexo 2. Goles local vs visitante",
        formula="GF/PJ = goles a favor / partidos jugados",
        explicacion=(
            "Este gráfico compara la producción ofensiva como local y visitante. Sirve para complementar "
            "el análisis de puntos con una variable deportiva distinta.",
        ),
    )

    documento.save(ARCHIVO_SALIDA)

    print("Informe Word generado correctamente en:")
    print(ARCHIVO_SALIDA)


if __name__ == "__main__":
    main()